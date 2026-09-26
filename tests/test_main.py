import json
import sys
import pytest
from openpyxl import load_workbook
from humble_catalog import db
from humble_catalog.__main__ import main

def test_bundle_command_prints_the_report(monkeypatch, tmp_path, capsys):
    monkeypatch.chdir(tmp_path)
    conn = db.connect("catalog.db")
    conn.execute("INSERT INTO items (machine_name, name, type) "
                 "VALUES ('owned_examplepress', 'Unrelated Book', 'ebook')")
    conn.commit()
    conn.close()
    fake = {
        "basic_data": {"human_name": "Bundle One", "currency": "USD"},
        "tier_item_data": {"owned_examplepress": {"human_name": "Unrelated Book"},
                           "new_examplepress": {"human_name": "The Hollow Crypt"}},
        "tier_display_data": {"initial": {
            "tier_item_machine_names": ["owned_examplepress",
                                        "new_examplepress"]}},
        "tier_pricing_data": {"initial": {
            "price|money": {"currency": "USD", "amount": 12.0}}},
    }
    from humble_catalog import bundle_preview
    monkeypatch.setattr(bundle_preview, "fetch_bundle",
                        lambda url, http=None: fake)
    monkeypatch.setattr(sys, "argv", [
        "humble_catalog", "bundle",
        "https://www.humblebundle.com/books/bundle-one-books"])
    main()
    out = capsys.readouterr().out
    assert "Bundle One" in out
    assert "owned 1" in out and "new 1" in out


def test_bundle_command_reports_a_bad_url_without_a_traceback(monkeypatch,
                                                              tmp_path, capsys):
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    monkeypatch.setattr(sys, "argv", [
        "humble_catalog", "bundle", "https://example.test/books/x"])
    with pytest.raises(SystemExit):
        main()
    assert "not a HumbleBundle URL" in capsys.readouterr().err


def test_export_command_writes_bom_csv(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()  # empty schema in cwd
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export"])
    main()
    out = (tmp_path / "catalog.csv").read_bytes()
    assert out.startswith(b"\xef\xbb\xbftitle,")  # UTF-8 BOM for Excel
    assert "Wrote 0 items to catalog.csv" in capsys.readouterr().out

def test_export_command_custom_path(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export", "my.csv"])
    main()
    assert (tmp_path / "my.csv").exists()
    assert "Wrote 0 items to my.csv" in capsys.readouterr().out

def test_export_command_writes_a_workbook_for_an_xlsx_path(tmp_path,
                                                           monkeypatch, capsys):
    # The suffix is the only format signal -- there is deliberately no
    # --format flag, since it could only duplicate or contradict the
    # filename beside it.
    monkeypatch.chdir(tmp_path)
    conn = db.connect("catalog.db")
    conn.execute("INSERT INTO items (machine_name, name, type) "
                 "VALUES ('m', 'A Quiet Life in Harbors', 'ebook')")
    conn.execute("INSERT INTO enrichment (item_id) VALUES (1)")
    conn.commit()
    conn.close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export", "my.xlsx"])
    main()
    ws = load_workbook(tmp_path / "my.xlsx").active
    assert ws.title == "Catalog"
    assert ws.cell(row=2, column=1).value == "A Quiet Life in Harbors"
    assert "Wrote 1 items to my.xlsx" in capsys.readouterr().out

def test_export_command_rejects_an_unknown_suffix(tmp_path, monkeypatch):
    # Guessing a format would write a mislabelled file.
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export", "my.txt"])
    with pytest.raises(SystemExit):
        main()
    assert not (tmp_path / "my.txt").exists()

def test_stats_command_prints_every_section_and_the_total(tmp_path, monkeypatch,
                                                          capsys):
    monkeypatch.chdir(tmp_path)
    conn = db.connect("catalog.db")
    # one item with no rating, no cover, no source_url -> a gap in all three
    conn.execute("INSERT INTO items (machine_name, name, type) "
                 "VALUES ('m', 'Unrelated Book', 'ebook')")
    conn.execute("INSERT INTO enrichment (item_id) VALUES (1)")
    conn.commit()
    conn.close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "stats"])
    main()
    out = capsys.readouterr().out
    for heading in ("By type", "Ratings", "Reading status",
                    "Enrichment", "Gaps", "Genres"):
        assert heading in out
    # the gaps section still names all three, as the gaps command did
    assert "Unrated" in out and "No cover" in out and "No source URL" in out
    # enrichment defaults to pending, and read_status to unread
    assert "1  Pending" in out and "1  Unread" in out
    assert "1  items total" in out


def test_import_sheets_dispatch(monkeypatch):
    calls = {}
    monkeypatch.setattr("humble_catalog.import_sheets.run",
                        lambda paths=None: calls.setdefault("paths", paths))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "import-sheets"])
    main()
    assert calls["paths"] is None

def test_import_sheets_dispatch_with_files(monkeypatch):
    calls = {}
    monkeypatch.setattr("humble_catalog.import_sheets.run",
                        lambda paths=None: calls.setdefault("paths", paths))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "import-sheets", "a.xlsx"])
    main()
    assert calls["paths"] == ["a.xlsx"]

def test_update_dispatch_passes_every_flag(monkeypatch):
    calls = {}
    monkeypatch.setattr("humble_catalog.update.run",
                        lambda **kw: calls.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "update",
                                      "--no-login", "--games", "--no-harvest"])
    main()
    assert calls == {"allow_login": False, "games": True, "no_harvest": True}


def test_update_dispatch_defaults(monkeypatch):
    calls = {}
    monkeypatch.setattr("humble_catalog.update.run",
                        lambda **kw: calls.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "update"])
    main()
    assert calls == {"allow_login": True, "games": False, "no_harvest": False}


def test_update_with_an_expired_session_says_so_without_a_traceback(
        monkeypatch):
    # The viewer's job panel offers its Log in button when the log says
    # "session expired" -- the same words extract's own exit uses.
    from humble_catalog import humble_api

    def logged_out(**kw):
        raise humble_api.NotLoggedIn("no session")
    monkeypatch.setattr("humble_catalog.update.run", logged_out)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "update"])
    with pytest.raises(SystemExit, match="session expired"):
        main()


def test_reset_dispatch(monkeypatch):
    calls = {}
    monkeypatch.setattr("humble_catalog.reset.run",
                        lambda: calls.setdefault("ran", True))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "reset"])
    main()
    assert calls["ran"] is True


def test_serve_uses_default_port(monkeypatch):
    seen = {}
    monkeypatch.setattr("humble_catalog.webapp.serve",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "serve"])
    main()
    assert seen == {"port": 8087, "terminal_commands": True}

def test_serve_accepts_a_port(monkeypatch):
    # the wrapper scripts honour HUMBLE_PORT, which is a lie unless the
    # port actually reaches the server
    seen = {}
    monkeypatch.setattr("humble_catalog.webapp.serve",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv",
                        ["humble_catalog", "serve", "--port", "8091"])
    main()
    assert seen == {"port": 8091, "terminal_commands": True}

def test_export_command_selects_columns(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    conn = db.connect("catalog.db")
    conn.execute("INSERT INTO items (machine_name, name, type) "
                 "VALUES ('m', 'A Quiet Life in Harbors', 'ebook')")
    conn.execute("INSERT INTO enrichment (item_id) VALUES (1)")
    conn.commit()
    conn.close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export", "my.csv",
                                      "--columns", "type, title"])
    main()
    lines = (tmp_path / "my.csv").read_text(encoding="utf-8-sig").splitlines()
    # whitespace tolerated; canonical order, not the requested order
    assert lines[0] == "title,type"
    assert lines[1] == "A Quiet Life in Harbors,ebook"

def test_export_command_rejects_an_unknown_column(tmp_path, monkeypatch,
                                                  capsys):
    # Loud here, silent on the web route: a typo on a command line is a
    # mistake being made now, not stored state outliving a rename.
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export", "my.csv",
                                      "--columns", "title,authorz"])
    with pytest.raises(SystemExit):
        main()
    assert "authorz" in capsys.readouterr().err
    assert not (tmp_path / "my.csv").exists()

def test_export_command_rejects_an_empty_column_list(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "export", "my.csv",
                                      "--columns", ""])
    with pytest.raises(SystemExit):
        main()
    assert not (tmp_path / "my.csv").exists()


# --- dependency guard ------------------------------------------------------

def test_no_dependencies_are_missing_in_a_working_install():
    from humble_catalog.__main__ import missing_dependencies
    assert missing_dependencies() == []


def test_the_guard_names_the_package_and_the_interpreter(monkeypatch, capsys):
    # The failure this replaces is a ModuleNotFoundError four imports deep,
    # which reads as a broken install. Naming sys.executable is what makes
    # "you are running the wrong python" visible.
    import importlib.util
    from humble_catalog import __main__ as entry

    real = importlib.util.find_spec
    monkeypatch.setattr(importlib.util, "find_spec",
                        lambda name: None if name == "rapidfuzz" else real(name))
    with pytest.raises(SystemExit) as exc:
        entry.check_dependencies()
    assert exc.value.code == 1
    err = capsys.readouterr().err
    assert "rapidfuzz" in err
    assert sys.executable in err
    assert "activate" in err.lower()


def test_the_guard_lists_every_missing_package_at_once(monkeypatch, capsys):
    import importlib.util
    from humble_catalog import __main__ as entry

    real = importlib.util.find_spec
    monkeypatch.setattr(
        importlib.util, "find_spec",
        lambda name: None if name in {"rapidfuzz", "flask"} else real(name))
    with pytest.raises(SystemExit):
        entry.check_dependencies()
    err = capsys.readouterr().err
    assert "flask" in err and "rapidfuzz" in err


def test_help_still_works_without_the_dependencies(monkeypatch, capsys):
    # The guard runs after parsing, so --help stays useful on a broken
    # environment -- it is stdlib-only and is how you find the command
    # names in the first place.
    import importlib.util
    real = importlib.util.find_spec
    monkeypatch.setattr(importlib.util, "find_spec", lambda name: None)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "--help"])
    with pytest.raises(SystemExit) as exc:
        main()
    assert exc.value.code == 0
    assert "usage: humble_catalog" in capsys.readouterr().out
    monkeypatch.setattr(importlib.util, "find_spec", real)


def test_the_guard_tracks_the_declared_dependencies():
    # Drift guard: a dependency added to pyproject but not here would be
    # invisible to the check, which is exactly the gap that produced the
    # confusing rapidfuzz traceback.
    import re
    import tomllib
    from pathlib import Path
    from humble_catalog.__main__ import RUNTIME_DEPENDENCIES

    data = tomllib.loads(
        (Path(__file__).parent.parent / "pyproject.toml").read_text("utf-8"))
    declared = {re.split(r"[<>=!~\[;\s]", d)[0]
                for d in data["project"]["dependencies"]}
    assert set(RUNTIME_DEPENDENCIES.values()) == declared


def test_harvest_dispatch_honours_the_stored_quota_by_default(monkeypatch):
    seen = {}
    monkeypatch.setattr("humble_catalog.harvest.run",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "harvest"])
    main()
    assert seen == {"ignore_quota": False}


def test_harvest_accepts_ignore_quota(monkeypatch):
    # the escape hatch is a lie unless the flag actually reaches the run
    seen = {}
    monkeypatch.setattr("humble_catalog.harvest.run",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv",
                        ["humble_catalog", "harvest", "--ignore-quota"])
    main()
    assert seen == {"ignore_quota": True}


def _seed_one_key(tmp_path, monkeypatch):
    """A catalog holding a single unactivated steam key."""
    monkeypatch.chdir(tmp_path)
    conn = db.connect("catalog.db")
    conn.execute("INSERT INTO bundles (gamekey, name, url, purchased_at) "
                 "VALUES ('kv789', 'Humble Game Bundle: Key Vault', "
                 "'https://example.invalid/kv789', '2024-01-02T00:00:00')")
    conn.execute(
        "INSERT INTO external_keys "
        "(gamekey, machine_name, human_name, key_type, raw) "
        "VALUES ('kv789', 'cindervale_ex', 'Cinder Vale', 'steam', ?)",
        (json.dumps({"human_name": "Cinder Vale", "key_type": "steam",
                     "machine_name": "cindervale_ex"}),))
    conn.commit()
    conn.close()


def test_keys_command_prints_the_summary(monkeypatch, tmp_path, capsys):
    _seed_one_key(tmp_path, monkeypatch)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "keys"])
    main()
    out = capsys.readouterr().out
    # "1 key", not "1 keys" -- a one-item tier read "1 items" in the bundle
    # preview, which no test caught and a browser did.
    assert "1 key -" in out
    # Undated, so the default report counts it without listing it.
    assert "Cinder Vale" not in out


def test_keys_all_lists_the_undated_rows(monkeypatch, tmp_path, capsys):
    _seed_one_key(tmp_path, monkeypatch)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "keys", "--all"])
    main()
    assert "Cinder Vale" in capsys.readouterr().out

def test_help_lists_the_pipeline_in_the_order_you_run_it(monkeypatch, capsys):
    # argparse prints subcommands in declaration order, and enrich matches
    # against what harvest cached - so listing enrich first told a
    # first-time reader to run the two backwards. Pinned because the
    # ordering is a property of where a parser happens to be declared,
    # which the next subcommand added could silently disturb.
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "--help"])
    with pytest.raises(SystemExit):
        main()
    out = capsys.readouterr().out
    assert out.index("\n    extract") < out.index("\n    harvest")
    assert out.index("\n    harvest") < out.index("\n    enrich")
    assert out.index("\n    enrich") < out.index("\n    serve")


def test_choice_command_prints_the_month(monkeypatch, tmp_path, capsys):
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    fake = {
        "baseSubscriptionPrice|money": {"currency": "USD", "amount": 11.99},
        "contentChoiceOptions": {
            "title": "January 2031",
            "contentChoiceState": {"initial": {"choices_made": []}},
            "contentChoiceData": {
                "extras": [],
                "game_data": {
                    "lanternlockpick_choice": {
                        "title": "Lantern & Lockpick",
                        "delivery_methods": ["steam"]}}}},
    }
    from humble_catalog import choice_preview, humble_api
    monkeypatch.setattr(humble_api, "ensure_login", lambda *a, **k: object())
    monkeypatch.setattr(choice_preview, "fetch_choice", lambda client=None: fake)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "choice"])
    main()
    out = capsys.readouterr().out
    assert "Humble Choice: January 2031" in out
    assert "new 1" in out


def test_choice_command_reports_a_dead_month_without_a_traceback(
        monkeypatch, tmp_path, capsys):
    monkeypatch.chdir(tmp_path)
    db.connect("catalog.db").close()
    from humble_catalog import choice_preview, humble_api

    def boom(client=None):
        raise ValueError("no Humble Choice month on offer")

    monkeypatch.setattr(humble_api, "ensure_login", lambda *a, **k: object())
    monkeypatch.setattr(choice_preview, "fetch_choice", boom)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "choice"])
    with pytest.raises(SystemExit):
        main()
    assert "no Humble Choice month on offer" in capsys.readouterr().err


def test_extract_no_login_reports_an_expired_session_without_a_traceback(
        monkeypatch, tmp_path, capsys):
    from humble_catalog import extract, humble_api

    monkeypatch.chdir(tmp_path)

    def fake_run(**kwargs):
        assert kwargs["allow_login"] is False
        raise humble_api.NotLoggedIn("expired")

    monkeypatch.setattr(extract, "run", fake_run)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "extract", "--no-login"])
    with pytest.raises(SystemExit) as exc:
        main()
    assert "login" in str(exc.value)


def test_serve_no_handoff_turns_the_terminal_commands_off(monkeypatch):
    # What the detached wrappers pass (#98): their console is hidden, so a
    # handed-over reset would wait for input nobody can type.
    seen = {}
    monkeypatch.setattr("humble_catalog.webapp.serve",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "serve",
                                      "--no-handoff"])
    main()
    assert seen == {"port": 8087, "terminal_commands": False}


def test_serve_lan_passes_no_handoff_too(monkeypatch):
    seen = {}
    monkeypatch.setattr("humble_catalog.webapp.serve",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "serve", "--lan",
                                      "--no-handoff"])
    main()
    assert seen["terminal_commands"] is False


def test_serve_lan_passes_its_options(monkeypatch):
    seen = {}
    monkeypatch.setattr("humble_catalog.webapp.serve",
                        lambda **kw: seen.update(kw))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "serve", "--lan",
                                      "--lan-host", "10.0.0.5", "--setup"])
    main()
    assert seen["port"] == 8087
    assert (seen["lan"].host, seen["lan"].port, seen["lan"].setup,
            seen["lan"].new_token) == ("10.0.0.5", None, True, False)


@pytest.mark.parametrize("flag", [["--lan-host", "10.0.0.5"],
                                  ["--lan-port", "9000"],
                                  ["--setup"], ["--new-token"]])
def test_lan_flags_need_lan(monkeypatch, capsys, flag):
    monkeypatch.setattr("humble_catalog.webapp.serve", lambda **kw: None)
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "serve", *flag])
    with pytest.raises(SystemExit):
        main()
    assert "needs --lan" in capsys.readouterr().err


def test_login_dispatches_to_humble_api_login(monkeypatch):
    from humble_catalog import humble_api
    called = []
    monkeypatch.setattr(humble_api, "login", lambda: called.append(True))
    monkeypatch.setattr(sys, "argv", ["humble_catalog", "login"])
    main()
    assert called == [True]


# -- A progress line can never be what kills a command (#102) -----------
# A terminal run redirected to a file (`enrich > log.txt`) writes with
# the locale code page, and one title it cannot encode used to end a
# long job with UnicodeEncodeError. harden_stdio() makes an unencodable
# character print as an escape instead.
def test_harden_stdio_prints_an_unencodable_character_as_an_escape():
    import os
    import subprocess
    env = {**os.environ, "PYTHONIOENCODING": "cp1252"}
    proc = subprocess.run(
        [sys.executable, "-c",
         "from humble_catalog.__main__ import harden_stdio; harden_stdio(); "
         "print('Tit\\u0307le')"],
        capture_output=True, env=env, timeout=30)
    assert proc.returncode == 0, proc.stderr.decode(errors="replace")
    assert proc.stdout.strip() == b"Tit\\u0307le"
